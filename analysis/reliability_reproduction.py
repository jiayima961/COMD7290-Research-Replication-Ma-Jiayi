"""Python translation of the identity-weighted, rater-conditional kappaetc suite.

Primary implementation: Daniel Klein, kappaetc 2.1.0 (11 Aug 2022), downloaded
from SSC. No Stata executable is used. The research script (1 Nov 2020) does not
pin its installed kappaetc build. Basic coefficients/SEs follow the recovered
source; benchmarking is explicitly versioned, because 2.1.0 changed it.
"""
from __future__ import annotations

import argparse
import ast
import contextlib
import hashlib
import importlib.metadata
import io
import json
import platform
import sys
import time
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from scipy.stats import t

SPEC_VERSION = "author-reliability-2020-11-01__identity-conditional__python-v1"
NAMES = ["percent_agreement", "brennan_prediger", "cohen_conger_kappa",
         "scott_fleiss", "gwet_ac1", "krippendorff_alpha"]
SOURCE_URL = "http://fmwww.bc.edu/repec/bocode/k/kappaetc.ado"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write_json(path: Path, obj) -> None:
    path.write_text(json.dumps(obj, indent=2, ensure_ascii=True, allow_nan=False,
                              default=lambda x: x.item() if isinstance(x, np.generic) else str(x)) + "\n",
                    encoding="utf-8")


def coefficients(raw: np.ndarray) -> tuple[list[dict], np.ndarray, dict]:
    """Translate kappaetc Mata lines 2555-2702, 2892-3092 for identity weights.

    Excel blanks are NaN, while -99 is a category, exactly as in author syntax.
    All-blank rows/columns are dropped. Single-rating rows inform expected
    agreement for the first five coefficients; alpha uses >=2-rating rows only.
    SE(alpha) deliberately uses the uncorrected alpha influence expression, as
    the recovered source does (not a guessed finite-sample delta correction).
    """
    raw = np.asarray(raw, dtype=float)
    if raw.ndim != 2 or raw.shape[1] < 2:
        raise ValueError("A subject-by-rater matrix with at least two slots is required.")
    if np.isinf(raw).any():
        raise ValueError("Infinite ratings are invalid.")
    x = raw[np.isfinite(raw).any(axis=1)]
    x = x[:, np.isfinite(x).any(axis=0)]
    if x.shape[1] < 2:
        raise ValueError("Fewer than two observed rating slots remain.")
    n, r = x.shape
    cat = np.unique(x[np.isfinite(x)])
    q = len(cat)
    if q < 2:
        raise ValueError("Ratings do not vary; chance-corrected coefficients are undefined.")
    observed = np.isfinite(x)
    ri = observed.sum(axis=1).astype(float)
    counts = np.stack([(x == k).sum(axis=1) for k in cat], axis=1).astype(float)
    prop = counts / ri[:, None]
    eligible = ri > 1
    npair = int(eligible.sum())
    if npair < 2:
        raise ValueError("Fewer than two subjects with at least two ratings.")
    pi = prop.mean(axis=0)
    ng = observed.sum(axis=0)
    pg = np.stack([(x == k).sum(axis=0) / ng for k in cat], axis=1)
    pair_sum = np.sum(counts * (counts - 1), axis=1)
    pai = np.divide(pair_sum, ri * (ri - 1), out=np.zeros(n), where=eligible)
    pa = float(pai[eligible].mean())
    peis = [np.zeros(n), np.full(n, 1 / q), np.zeros(n), prop @ pi,
            prop @ (1 - pi) / (q - 1)]
    # Identity weights simplify the source Kronecker products to these sums.
    for kidx, k in enumerate(cat):
        delta = (x == k).astype(float)
        adjusted = delta - pg[None, :, kidx] * (observed - ng[None, :] / n)
        other_marginals = pg[:, kidx].sum() - pg[:, kidx]
        peis[2] += np.sum(adjusted * (n / ng)[None, :] * other_marginals[None, :], axis=1)
    peis[2] /= r * (r - 1)
    rbar = float(ri[eligible].mean())
    alpha_pa_i = np.divide(pair_sum, rbar * (ri - 1), out=np.zeros(n), where=eligible)
    alpha_pa = float(alpha_pa_i[eligible].mean())
    alpha_pa_i = (alpha_pa_i - alpha_pa * (ri - rbar) / rbar) * eligible
    alpha_pi = counts[eligible].sum(axis=0) / (npair * rbar)
    alpha_pe_i = counts @ alpha_pi / rbar
    alpha_pe = float(alpha_pe_i[eligible].mean())
    alpha_pe_i = (alpha_pe_i - alpha_pe * (ri - rbar) / rbar) * eligible
    peis.append(alpha_pe_i)
    result, influences = [], []
    for j, name in enumerate(NAMES):
        nn = npair if j == 5 else n
        pe_i = peis[j]
        pe = float(pe_i[eligible].mean() if j == 5 else pe_i.mean())
        this_pa = alpha_pa if j == 5 else pa
        this_pai = alpha_pa_i if j == 5 else pai
        eps = 1 / (npair * rbar) if j == 5 else 0.0
        bprime = (this_pa - pe) / (1 - pe)
        estimate = ((1 - eps) * this_pa + eps - pe) / (1 - pe)
        bistar = nn / npair * (this_pai - pe) / (1 - pe) * eligible
        correction = 2 * (1 - bprime) * (pe_i - pe) / (1 - pe)
        bistar -= correction * (eligible if j == 5 else 1)
        influence = (bistar - bprime) * (eligible if j == 5 else 1)
        variance = float(np.sum(influence ** 2) / (nn * (nn - 1)))
        se = float(np.sqrt(variance))
        df = nn - 1
        critical = float(t.ppf(.975, df))
        result.append({"coefficient": name, "estimate": float(estimate), "se": se,
                       "ci95_low": max(-1., float(estimate - critical * se)),
                       "ci95_high": min(1., float(estimate + critical * se)),
                       "ci95_low_unclipped": float(estimate - critical * se),
                       "ci95_high_unclipped": float(estimate + critical * se),
                       "df": df, "t": float(estimate / se) if se else None,
                       "pvalue": float(2 * t.sf(abs(estimate / se), df)) if se else None,
                       "observed_agreement": float((1 - eps) * this_pa + eps),
                       "expected_agreement": pe, "n_expected": nn,
                       "n_observed": npair, "n_rating_slots": r, "n_categories": q,
                       "scale": "coefficient", "se_type": "conditional_on_rating_slots",
                       "interval_distribution": "Student_t", "weights": "identity",
                       "source_implementation": "kappaetc_2.1.0_11aug2022_Mata_translation",
                       "status": "computed_from_author_matrix__original_build_unpinned"})
        influences.append(influence)
    audit = {"input_rows": len(raw), "input_columns": raw.shape[1], "n_any_rating": n,
             "n_at_least_two": npair, "all_blank_rows_excluded": len(raw) - n,
             "all_blank_columns_excluded": raw.shape[1] - r,
             "ratings_per_subject_counts": {int(k): int(v) for k, v in sorted(Counter(ri.astype(int)).items())},
             "categories_inferred_from_observed_values": cat.tolist(),
             "numeric_minus99_is_retained_category": bool((-99 == cat).any()),
             "slot_marginals": pg.tolist(), "subject_averaged_category_proportions": pi.tolist(),
             "alpha_pooled_proportions": alpha_pi.tolist(), "alpha_epsilon": 1 / (npair * rbar)}
    return result, np.stack(influences, axis=1), audit


def benchmark(row: dict, mode: str) -> list[dict]:
    """Landis-Koch, descending, strict cumulative probability >95%.

    current_truncated is a direct translation of recovered 2.1.0 lines 4689-4733.
    legacy_untruncated_reconstruction implements the author's documented pre-
    2022 behavior: no rescaling, cumulative mass forced to 1 in the last row.
    The original 2.0.0 source has not been recovered, so this stays labelled.
    """
    b, se, df = row["estimate"], row["se"], row["df"]
    lb = 0. if row["coefficient"] == "percent_agreement" else -1.
    upper = np.array([1., .8, .6, .4, .2, 0.])
    bounds = np.r_[upper, lb]
    dist = t.cdf((b - bounds) / se, df)
    membership = np.diff(dist)
    trunc = float(dist[-1] - dist[0])
    if mode == "current_truncated":
        membership = membership / trunc
    cumulative = np.cumsum(membership)
    if mode == "legacy_untruncated_reconstruction":
        # Pre-2022 source substitutes probability 1 for the missing lower limit.
        # The primary author's 2020 explanation explicitly states that the
        # final cumulative probability is forced to 1. Allocate the residual
        # to that interval; this is labelled a reconstruction, not recovered code.
        membership[-1] = 1 - membership[:-1].sum()
        cumulative = np.cumsum(membership)
    candidates = np.flatnonzero(cumulative > .95)
    chosen = int(candidates[0]) if len(candidates) else None
    labels = ["Almost Perfect", "Substantial", "Moderate", "Fair", "Slight", "Poor"]
    out = []
    for i in range(6):
        out.append({"coefficient": row["coefficient"], "benchmark_mode": mode,
                    "benchmark": labels[i], "lower": (None if mode == "legacy_untruncated_reconstruction" and i == 5 else float(bounds[i + 1])), "upper": float(upper[i]),
                    "interval_membership_probability": float(membership[i]),
                    "cumulative_probability_from_highest": float(cumulative[i]),
                    "selected_at_95_percent": i == chosen,
                    "truncation_mass": trunc if mode == "current_truncated" else None,
                    "distribution": "Student_t", "df": df,
                    "status": "source_verified_2.1.0" if mode == "current_truncated"
                    else "pre_2022_behavior_reconstruction__original_build_unpinned"})
    return out


def independent_weighted_points(raw: np.ndarray, weights: np.ndarray, anchor_pair_fraction: bool = True) -> np.ndarray:
    """Independent direct-pair implementation for numerical derivative checking.

    This intentionally uses direct comparisons and independently defined margins,
    rather than the analytical influence expressions in coefficients(). Alpha
    is returned before its finite-sample epsilon correction, matching source V().
    For ordinary measures, the source's IF anchors nprime/n at its observed value
    instead of differentiating that ratio. This distinction matters with singleton
    rows and is made explicit here. Disabling the anchor yields the ordinary
    weighted point estimator, which is not the source's variance functional.
    """
    x = np.asarray(raw, float)
    keep = np.isfinite(x).any(axis=1)
    x, w = x[keep], weights[keep]
    x = x[:, np.isfinite(x).any(axis=0)]
    cats = np.unique(x[np.isfinite(x)])
    npresent = np.isfinite(x).sum(axis=1)
    pair_rows = npresent > 1
    agreement = []
    fractions = []
    for subject in x:
        vals = subject[np.isfinite(subject)]
        fractions.append([np.mean(vals == k) for k in cats])
        pairs = [float(vals[a] == vals[b]) for a in range(len(vals)) for b in range(a)]
        agreement.append(np.mean(pairs) if pairs else 0.)
    agreement, fractions = np.array(agreement), np.array(fractions)
    pa = np.average(agreement[pair_rows], weights=w[pair_rows])
    pi = np.average(fractions, axis=0, weights=w)
    slot_margin = []
    for slot in x.T:
        mask = np.isfinite(slot)
        slot_margin.append([np.average(slot[mask] == k, weights=w[mask]) for k in cats])
    pair_pe = [np.dot(slot_margin[a], slot_margin[b]) for a in range(len(slot_margin)) for b in range(a)]
    pe = [0., 1 / len(cats), np.mean(pair_pe), np.dot(pi, pi),
          (1 - np.dot(pi, pi)) / (len(cats) - 1)]
    if anchor_pair_fraction:
        fraction = pair_rows.mean()
        values = [np.average(pair_rows * (agreement - p), weights=w) / fraction / (1 - p) for p in pe]
    else:
        values = [(pa - p) / (1 - p) for p in pe]
    rating_weights = w[pair_rows] * npresent[pair_rows]
    alpha_pa = np.average(agreement[pair_rows], weights=rating_weights)
    alpha_pi = np.average(fractions[pair_rows], axis=0, weights=rating_weights)
    alpha_pe = np.dot(alpha_pi, alpha_pi)
    values.append((alpha_pa - alpha_pe) / (1 - alpha_pe))
    return np.array(values)


def tests() -> dict:
    import krippendorff
    from statsmodels.stats.inter_rater import cohens_kappa
    table = np.array([[20, 5], [10, 15]])
    data = np.array([[i, j] for i in range(2) for j in range(2) for _ in range(table[i, j])], float)
    rows, inf, _ = coefficients(data)
    actual = np.array([r["estimate"] for r in rows])
    expected = np.array([.7, .4, .4, 13 / 33, 41 / 101, .4])
    np.testing.assert_allclose(actual, expected, rtol=0, atol=1e-12)
    np.testing.assert_allclose(rows[0]["se"], np.sqrt(.7 * .3 / 49), atol=1e-14)
    np.testing.assert_allclose(rows[1]["se"], 2 * np.sqrt(.7 * .3 / 49), atol=1e-14)
    np.testing.assert_allclose(actual[2], cohens_kappa(table).kappa, atol=1e-14)
    np.testing.assert_allclose(actual[5], krippendorff.alpha(data.T, level_of_measurement="nominal"), atol=1e-14)
    missing = np.array([[0, 0, np.nan, 1], [1, 1, 1, np.nan], [0, 1, 0, 0],
                        [1, np.nan, np.nan, np.nan], [np.nan, np.nan, np.nan, np.nan],
                        [0, 1, np.nan, np.nan], [np.nan, 1, 0, np.nan]], float)
    mrows, minf, maudit = coefficients(missing)
    np.testing.assert_allclose(mrows[5]["estimate"], krippendorff.alpha(missing.T, level_of_measurement="nominal"), atol=1e-14)
    mclean = missing[np.isfinite(missing).any(axis=1)]
    derivative_error = 0.
    unanchored_difference = 0.
    h = 1e-5
    for i in range(len(mclean)):
        wp, wm = np.ones(len(mclean)), np.ones(len(mclean))
        wp[i] += h
        wm[i] -= h
        derivative = (independent_weighted_points(mclean, wp) - independent_weighted_points(mclean, wm)) / (2 * h)
        expected_if = derivative * np.array([len(mclean)] * 5 + [maudit["n_at_least_two"]])
        derivative_error = max(derivative_error, float(np.max(np.abs(expected_if - minf[i]))))
        ordinary_derivative = (independent_weighted_points(mclean, wp, False) - independent_weighted_points(mclean, wm, False)) / (2 * h)
        ordinary_if = ordinary_derivative * np.array([len(mclean)] * 5 + [maudit["n_at_least_two"]])
        unanchored_difference = max(unanchored_difference, float(np.max(np.abs(ordinary_if - minf[i]))))
    assert derivative_error < 1e-8, derivative_error
    np.testing.assert_allclose(inf.mean(axis=0), 0, atol=1e-14)
    bad = data.copy()
    bad[0, 0] = np.inf
    try:
        coefficients(bad)
        raise AssertionError("Negative test failed to reject infinity.")
    except ValueError as e:
        detected = str(e)
    recovered = coefficients(data)[0][2]["estimate"]
    assert abs(recovered - .4) < 1e-12
    benchmark_example = {"coefficient": "gwet_ac1", "estimate": .97, "se": .02, "df": 100}
    bm_new = benchmark(benchmark_example, "current_truncated")
    bm_old = benchmark(benchmark_example, "legacy_untruncated_reconstruction")
    assert [b["benchmark"] for b in bm_new if b["selected_at_95_percent"]] == ["Almost Perfect"]
    assert [b["benchmark"] for b in bm_old if b["selected_at_95_percent"]] == ["Poor"]
    assert abs(sum(b["interval_membership_probability"] for b in bm_new) - 1) < 1e-14
    assert abs(sum(b["interval_membership_probability"] for b in bm_old) - 1) < 1e-14
    return {"status": "passed", "synthetic_known_answer_table": table.tolist(),
            "known_coefficients": dict(zip(NAMES, actual.tolist())),
            "percent_agreement_se_closed_form": rows[0]["se"],
            "independent_packages": ["statsmodels.cohens_kappa", "krippendorff.alpha(nominal)"],
            "missing_rater_alpha_independently_checked": mrows[5]["estimate"],
            "conditional_influence_finite_difference_max_abs_error": derivative_error,
            "ordinary_unanchored_point_derivative_differs_from_source_IF_max": unanchored_difference,
            "source_IF_validation_functional": "Observed >=2-rating fraction fixed; matches literal source b_istar. Ordinary weighted point-estimator derivative differs with singleton rows and is not asserted equal.",
            "benchmark_version_behavior_check": {"synthetic_b": .97, "se": .02, "df": 100,
                 "current_truncated_selection": "Almost Perfect", "legacy_reconstructed_selection": "Poor",
                 "interval_probability_sums_equal_one": True},
            "negative_test": {"input": "synthetic infinite rating", "detected": detected,
                              "recovery": "restored synthetic input; kappa returned to 0.4"}}


def assert_numeric(df: pd.DataFrame, cols: list[str], allowed: set[float]) -> np.ndarray:
    values = df[cols].apply(pd.to_numeric, errors="raise").to_numpy(dtype=float)
    observed = set(values[np.isfinite(values)].tolist())
    if not observed <= allowed:
        raise ValueError(f"Unexpected rating categories in {cols}: {sorted(observed - allowed)}")
    return values


def audit_matching(tone: pd.DataFrame, actor: pd.DataFrame) -> tuple[dict, pd.DataFrame]:
    rows = []
    for idx, row in actor.iterrows():
        coders = [str(row[f"Coder_{i}"]) for i in range(14) if pd.notna(row[f"Coder_{i}"])]
        keys = [str(row[f"key_{i}"]) for i in range(14) if pd.notna(row[f"key_{i}"])]
        key_text_mismatch = 0
        unparseable = 0
        for key in keys:
            try:
                key_text_mismatch += ast.literal_eval(key)[0] != row["text_id"]
            except (ValueError, SyntaxError, TypeError, IndexError):
                unparseable += 1
        rows.append({"excel_row": int(idx + 2), "text_id": row["text_id"],
                     "n_named_slots": len(coders), "n_distinct_named_coders": len(set(coders)),
                     "repeated_same_coder_in_row": len(coders) != len(set(coders)),
                     "n_nonmissing_keys": len(keys), "n_distinct_matching_keys": len(set(keys)),
                     "key_text_id_mismatches": key_text_mismatch, "unparseable_keys": unparseable})
    detail = pd.DataFrame(rows)
    audit = {"tone_rows": len(tone), "tone_unique_text_ids": int(tone["Text_ID"].nunique()),
             "tone_duplicate_text_ids": int(tone["Text_ID"].duplicated().sum()),
             "actor_rows": len(actor), "actor_unique_text_ids": int(actor["text_id"].nunique()),
             "actor_rows_missing_text_id": int(actor["text_id"].isna().sum()),
             "actor_duplicate_first_slot_keys": int(actor["key_0"].duplicated().sum()),
             "rows_with_multiple_distinct_keys": int((detail.n_distinct_matching_keys > 1).sum()),
             "key_text_id_mismatch_cells": int(detail.key_text_id_mismatches.sum()),
             "unparseable_key_cells": int(detail.unparseable_keys.sum()),
             "rows_with_repeated_same_named_coder": int(detail.repeated_same_coder_in_row.sum()),
             "rows_with_at_least_two_distinct_named_coders": int((detail.n_distinct_named_coders > 1).sum()),
             "actor_text_ids_absent_from_tone_workbook": int(len(set(actor.text_id) - set(tone.Text_ID))),
             "interpretation": "Faithful arithmetic on the author's matched workbook is reported. Positional slots are not stable named raters; repeated named coders raise independence concerns. No rematching or silent exclusions were performed."}
    return audit, detail


def main(project_root: Path, output_root: Path | None = None, *, diagnostics_root: Path | None = None, log_root: Path | None = None, prepared_root: Path | None = None, methods_root: Path | None = None, published_root: Path | None = None) -> None:
    start = time.perf_counter()
    project_root = project_root.resolve()
    output_root = (output_root or project_root / 'outputs/reproduced/reliability').resolve()
    diagnostics_root = (diagnostics_root or project_root / 'outputs/diagnostics/reliability').resolve()
    log_root = (log_root or project_root / 'records/run_logs/skill2/reliability').resolve()
    prepared_root = (prepared_root or project_root / 'data/prepared/reliability').resolve()
    methods_root = (methods_root or project_root / 'sources/local/skill2_methods/reliability').resolve()
    published_root = (published_root or project_root / 'outputs/published/reliability').resolve()
    for folder in [output_root, diagnostics_root, log_root, prepared_root, published_root]:
        folder.mkdir(parents=True, exist_ok=True)
    for required in ['kappaetc_2.1.0.ado', 'kappaetc_2.1.0.sthlp']:
        if not (methods_root / required).is_file():
            raise FileNotFoundError(f'Required archived primary method source is missing: {methods_root / required}')
    log = []
    def say(msg):
        print(msg, flush=True)
        log.append(msg)
    say("Applying $reproduce-and-diagnose: independent reliability forward test; Python only.")
    test_result = tests()
    write_json(diagnostics_root / "checks.json", test_result)
    say("Known-answer, independent alpha/Cohen, finite-difference influence, negative-test and recovery checks passed.")
    paths = {"ratings_tone": project_root / "data/original/independent_ratings_tone.xlsx",
             "ratings_actor_reference": project_root / "data/original/independent_ratings_actor_reference.xlsx",
             "author_reliability_script": project_root / "sources/author_code/reliability_original.do",
             "source_manifest": project_root / "records/source_manifest.csv",
             "target_article": project_root / "sources/local/Lob_et_al_2024.pdf",
             "invoked_skill": project_root / ".agents/skills/reproduce-and-diagnose/SKILL.md",
             "execution_contract": project_root / ".agents/skills/reproduce-and-diagnose/references/execution_contract.md"}
    inputs = {k: {"path": str(p), "sha256_before": sha256(p)} for k, p in paths.items()}
    registered = pd.read_csv(paths["source_manifest"], encoding="utf-8-sig").set_index("source_id")
    for ours, sid in [("ratings_tone", "ratings_tone"), ("ratings_actor_reference", "ratings_actref"),
                      ("author_reliability_script", "reliability_script"), ("target_article", "selected_article")]:
        assert inputs[ours]["sha256_before"] == registered.loc[sid, "sha256"], sid
        inputs[ours]["registered_manifest_hash_matches"] = True
    wb_t = load_workbook(paths["ratings_tone"], read_only=True, data_only=True)
    wb_a = load_workbook(paths["ratings_actor_reference"], read_only=True, data_only=True)
    assert wb_t.sheetnames == ["ton_codes"]
    assert wb_a.sheetnames == ["actref_codes_matched_num"]
    st, sa = wb_t["ton_codes"], wb_a["actref_codes_matched_num"]
    assert st.max_row == 1701 and st.max_column == 10
    assert sa.max_row == 3865 and sa.max_column == 225
    ht, ha = [c.value for c in next(st.iter_rows())], [c.value for c in next(sa.iter_rows())]
    tone = pd.read_excel(paths["ratings_tone"], sheet_name="ton_codes", usecols="A:J", nrows=1700)
    actor = pd.read_excel(paths["ratings_actor_reference"], sheet_name="actref_codes_matched_num", usecols="A:HQ", nrows=3864)
    assert len(tone) == 1700 and len(actor) == 3864
    responsiveness = ["GENUINE-KONSTRUIERTE_RESP_AKTEUR_NUM_" + str(i) for i in range(5)]
    placeholders = ["Q", "AG", "AW", "BM", "CC"]
    mapping = []
    for alias, header in zip(placeholders, responsiveness):
        col = ha.index(header) + 1
        assert get_column_letter(col) == alias
        assert len(header) > 32
        mapping.append({"stata_placeholder_in_author_script": alias, "excel_column": alias,
                        "excel_index_one_based": col, "original_header": header,
                        "header_length": len(header),
                        "evidence": "Author script explicitly identifies importer placeholders; workbook column letter and exact long family header independently agree."})
    write_json(diagnostics_root / "responsiveness_mapping.json", mapping)
    match_audit, match_detail = audit_matching(tone, actor)
    write_json(diagnostics_root / "matching_audit.json", match_audit)
    match_detail.to_csv(diagnostics_root / "matching_row_audit.csv", index=False)
    definitions = {"tone": (tone, [f"Ton_{i}" for i in range(1, 5)], {1., 2., 3., 4., 5.}),
                   "valence": (actor, [f"VALENZ_AKTEUR_NUM_{i}" for i in range(14)], {-99., 0., 1., 2.}),
                   "recognition": (actor, [f"ANERKENNUNG_AKTEUR_NUM_{i}" for i in range(14)], {0., 1.}),
                   "outrage": (actor, [f"HETZE_AKTEUR_NUM_{i}" for i in range(14)], {0., 1.}),
                   "responsiveness": (actor, responsiveness, {-99., 1., 2.})}
    results, bmarks, col_audit, samples, independent = [], [], [], {}, []
    for outcome, (frame, columns, allowed) in definitions.items():
        assert len(set(columns)) == len(columns)
        for c in columns:
            position = list(frame.columns).index(c) + 1
            counts = frame[c].value_counts(dropna=False)
            col_audit.append({"outcome": outcome, "original_column": c,
                              "excel_column": get_column_letter(position), "nonmissing": int(frame[c].notna().sum()),
                              "blank_missing": int(frame[c].isna().sum()), "numeric_minus99": int((frame[c] == -99).sum()),
                              "observed_values": ";".join(str(x) for x in sorted(frame[c].dropna().unique()))})
        for selection, selected in [("two_slots", columns[:2]), ("all_author_selected_slots", columns)]:
            ident = outcome + "__" + selection
            matrix = assert_numeric(frame, selected, allowed)
            rows, influence, audit = coefficients(matrix)
            audit["selected_original_columns"] = selected
            audit["input_domain_preserved"] = True
            samples[ident] = audit
            for row in rows:
                row.update({"analysis_id": ident, "outcome": outcome, "selection": selection,
                            "source_data_version": inputs["ratings_tone" if outcome == "tone" else "ratings_actor_reference"]["sha256_before"],
                            "specification_version": SPEC_VERSION})
                results.append(row)
                for mode in ["current_truncated", "legacy_untruncated_reconstruction"]:
                    for bm in benchmark(row, mode):
                        bm.update({"analysis_id": ident, "outcome": outcome, "selection": selection})
                        bmarks.append(bm)
            np.save(diagnostics_root / (ident + "__conditional_influence.npy"), influence)
            pd.DataFrame(matrix, columns=selected).assign(excel_row=np.arange(len(matrix)) + 2).to_csv(prepared_root / (ident + "__rating_matrix.csv"), index=False)
            # Independent alpha implementation on actual study inputs, not saved targets.
            import krippendorff
            alpha_check = float(krippendorff.alpha(matrix.T, level_of_measurement="nominal"))
            error = abs(alpha_check - rows[5]["estimate"])
            assert error < 1e-12, (ident, alpha_check, rows[5]["estimate"])
            independent.append({"analysis_id": ident, "coefficient": "krippendorff_alpha",
                                "independent_implementation": "krippendorff.alpha_nominal_0.9.0",
                                "independent_estimate": alpha_check, "translated_estimate": rows[5]["estimate"],
                                "absolute_difference": error, "passed": True})
            # Check analytical IF against numerical weighted functional derivatives on actual data.
            x = matrix[np.isfinite(matrix).any(axis=1)]
            points = sorted(set([0, len(x) // 3, len(x) // 2, len(x) - 1]))
            max_error = 0.
            for point in points:
                wp, wm = np.ones(len(x)), np.ones(len(x))
                h = 1e-4
                wp[point] += h
                wm[point] -= h
                deriv = (independent_weighted_points(x, wp) - independent_weighted_points(x, wm)) / (2 * h)
                expected = deriv * np.array([len(x)] * 5 + [audit["n_at_least_two"]])
                max_error = max(max_error, float(np.max(np.abs(influence[point] - expected))))
            assert max_error < 2e-7, (ident, max_error)
            independent.append({"analysis_id": ident, "coefficient": "all_six_conditional_influences",
                                "independent_implementation": "direct_pair_weighted_functional_central_differences",
                                "points_checked": points, "max_absolute_difference": max_error, "passed": True})
            say(f"Computed {ident}: N(any)={audit['n_any_rating']}, N(>=2)={audit['n_at_least_two']}, slots={audit['input_columns']}, alpha={rows[5]['estimate']:.10f}")
    pd.DataFrame(results).to_csv(output_root / "reliability_coefficients.csv", index=False)
    pd.DataFrame(bmarks).to_csv(output_root / "probabilistic_benchmarks.csv", index=False)
    pd.DataFrame(col_audit).to_csv(diagnostics_root / "column_audit.csv", index=False)
    write_json(diagnostics_root / "sample_audit.json", samples)
    write_json(diagnostics_root / "independent_study_checks.json", independent)
    for key, path in paths.items():
        inputs[key]["sha256_after"] = sha256(path)
        assert inputs[key]["sha256_before"] == inputs[key]["sha256_after"], key
    write_json(log_root / "input_manifest.json", inputs)
    versions = {p: importlib.metadata.version(p) for p in ["numpy", "pandas", "scipy", "openpyxl", "statsmodels", "krippendorff"]}
    versions.update({"python": sys.version, "platform": platform.platform()})
    write_json(log_root / "versions.json", versions)
    source_dir = methods_root
    primary = []
    for f in ["kappaetc_2.1.0.ado", "kappaetc_2.1.0.sthlp"]:
        p = source_dir / f
        if p.exists():
            primary.append({"file": str(p), "sha256": sha256(p), "bytes": p.stat().st_size})
    write_json(log_root / "method_provenance.json", {"primary_source_url": SOURCE_URL,
               "source_download_date": "2026-10-18", "recovered_version": "2.1.0 11aug2022",
               "source_files": primary, "translated_source_lines": {"raw_missing_handling": "2555-2702",
               "coefficients_and_conditional_SE": "2457-2477;2892-3092", "Student_t_CI": "4561-4617",
               "truncated_benchmark": "4675-4753", "clipped_CI": "5010-5021"},
               "legacy_benchmark_source": "Klein's 6 Dec 2020 explanation and 11 Aug 2022 version announcement; recovered legacy source when available",
               "legacy_benchmark_url": "https://www.statalist.org/forums/forum/general-stata-discussion/general/1368285-kappaetc-available-from-ssc/page3",
               "original_author_script_date": "2020-11-01", "original_installed_build": "unknown",
               "default_categories": "observed distinct numeric ratings; author specifies no categories()",
               "missing": "blank cells missing; -99 retained as observed numeric category",
               "uncertainty": "SE conditional on positional slots; Student t df N-1 (alpha N>=2-1), 95%; CI clipped to [-1,1]",
               "original_build_equivalence": "not established; no Stata was executed"})
    unknowns = ["The original installed kappaetc build was not pinned. Basic methods follow recovered 2.1.0 source; its changelog identifies only benchmarking as the change from 2.0.0.",
                "The script predates the 2022 truncated benchmark change. Both source-version benchmark methods are retained; exact historical installed build remains unknown.",
                "Rater variables are positional slots containing different humans across subjects. Some actor rows repeat one human in multiple slots. The arithmetic faithfully reruns the supplied author matrix; independence and the original matching algorithm are not established by this rerun.",
                "Numeric -99 is treated as a category because the original reliability script never declares it missing; interpreting -99 as missing would change the source specification.",
                "Responsiveness uses only first five original long-header fields, independently resolved against Excel letters Q/AG/AW/BM/CC and the author comments; other actor outcomes use all 14 fields.",
                "Conditional SEs follow the recovered source exactly. With singleton rows, its influence expression anchors the observed fraction with at least two ratings; it differs from differentiating an ordinary weighted point estimator that lets that fraction change. The numerical derivative check explicitly uses the source's anchored functional. Alpha SE uses the source's uncorrected-alpha influence, while its point coefficient includes epsilon correction.",
                "No published target CSV or previous AI reliability conclusion was used for implementation or fitting. No exact cross-software equivalence claim is made."]
    write_json(diagnostics_root / "unknowns_and_discrepancies.json", unknowns)
    elapsed = time.perf_counter() - start
    receipt = {"skill": "$reproduce-and-diagnose", "skill_path": str(paths["invoked_skill"]),
               "skill_sha256": inputs["invoked_skill"]["sha256_before"],
               "task": "independent primary-source forward test of real reliability suite", "authorization": "AGENTS D011 and parent task", "statistical_runtime": "Python only",
               "project_root": str(project_root), "output_root": str(output_root), "diagnostics_root": str(diagnostics_root), "log_root": str(log_root), "prepared_root": str(prepared_root), "methods_root": str(methods_root), "published_root": str(published_root), "specification_version": SPEC_VERSION,
               "command_argv": sys.argv, "timestamp_utc": datetime.now(timezone.utc).isoformat(),
               "runtime_seconds": elapsed, "module_sha256": sha256(Path(__file__)),
               "analyses_completed": 10, "coefficient_rows": len(results), "benchmark_rows": len(bmarks),
               "original_files_unchanged": True, "checks_passed": True,
               "no_Stata_execution": True, "new_dependencies_installed": False}
    write_json(log_root / "invocation_receipt.json", receipt)
    say(f"Completed 10 real analyses / {len(results)} coefficient rows / {len(bmarks)} benchmark rows in {elapsed:.2f}s; source inputs unchanged.")
    (log_root / "stdout.log").write_text("\n".join(log) + "\n", encoding="utf-8")
    table = pd.DataFrame(results).pivot(index=["outcome", "selection"], columns="coefficient", values="estimate").to_string(float_format=lambda z:f"{z:.6f}")
    report = "# Independent reliability forward test\n\nComputed from original independent-rating workbooks and the original author script. All computations ran in Python; no Stata executable or new installation was used.\n\n```text\n" + table + "\n```\n\n## Interpretation and limits\n\n" + "\n\n".join(unknowns) + "\n\n## Outputs\n\n- reliability_coefficients.csv: all six measures, conditional SEs, t tests, clipped and unclipped 95% intervals, expected/observed agreement, actual samples and source versions.\n- probabilistic_benchmarks.csv: all Landis-Koch interval probabilities and selected categories for separately identified pre-2022 and current methods.\n- sample_audit.json, column_audit.csv, matching_audit.json, matching_row_audit.csv and responsiveness_mapping.json: selection, missingness, numeric -99, exact headers, keys and named-coder audit.\n- checks.json and independent_study_checks.json: closed-form known answers, independent installed implementations, numerical influence checks and harmless synthetic negative test/recovery.\n- input_manifest.json, method_provenance.json, invocation_receipt.json, versions.json and stdout.log: exact input hashes, settings, source/code evidence and actual execution.\n- Rating matrices and conditional influence NPY files retain the calculation inputs and subject-level variance contributions.\n\nRun from the project root with `rtk proxy .venv/Scripts/python.exe -X utf8 analysis/reliability_reproduction.py --project-root .`. This reruns the statistical calculations and post-fit comparison offline from the archived primary page.\n"
    report += f"\n\nArtifact locations: results `{output_root}`; diagnostics `{diagnostics_root}`; run logs `{log_root}`; prepared rating matrices `{prepared_root}`; primary method evidence `{methods_root}`.\n"
    (output_root / "REPORT.md").write_text(report, encoding="utf-8")
    snapshot = methods_root / "official_reliabilities_2026-10-18.html"
    if snapshot.is_file():
        if __package__:
            from .reliability_primary_comparison import main as compare_primary
        else:
            from reliability_primary_comparison import main as compare_primary
        comparison_stdout = io.StringIO()
        with contextlib.redirect_stdout(comparison_stdout):
            compare_primary(snapshot, output_root, diagnostics_root, published_root, log_root)
        for line in comparison_stdout.getvalue().splitlines():
            say(line)
        with (output_root / "REPORT.md").open("a", encoding="utf-8") as stream:
            stream.write("\nAll 30 first/second-slot coefficients and all five sample counts match the archived primary website at displayed precision. Primary targets are saved separately from computed estimates. The website does not supply SE/interval or all-slot targets.\n")
    (log_root / "stdout.log").write_text("\n".join(log) + "\n", encoding="utf-8")
    files = [Path(__file__), Path(__file__).with_name("reliability_primary_comparison.py")]
    for folder in [output_root, diagnostics_root, log_root, prepared_root, methods_root, published_root]:
        files.extend(f for f in folder.rglob("*") if f.is_file() and f.name != "artifact_manifest.json")
    write_json(log_root / "artifact_manifest.json", [{"path": str(f), "sha256": sha256(f), "bytes": f.stat().st_size} for f in sorted(set(files))])


run = main


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project-root", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--output-root", type=Path)
    parser.add_argument("--diagnostics-root", type=Path)
    parser.add_argument("--log-root", type=Path)
    parser.add_argument("--prepared-root", type=Path)
    parser.add_argument("--methods-root", type=Path)
    parser.add_argument("--published-root", type=Path)
    args = parser.parse_args()
    main(args.project_root, args.output_root, diagnostics_root=args.diagnostics_root, log_root=args.log_root, prepared_root=args.prepared_root, methods_root=args.methods_root, published_root=args.published_root)
