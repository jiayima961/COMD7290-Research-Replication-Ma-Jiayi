"""Build the two English technical DOCX reports from their ordered Markdown sources.

Run with the bundled document runtime. This authoring step does not certify
rendering; the separate Word/PDF/PNG QA receipt binds the delivered bytes.
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
def build(source):
 d=Document();sec=d.sections[0]
 sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
 sec.page_width=Inches(8.5);sec.page_height=Inches(11)
 for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
  s=d.styles[name];s.font.name='Times New Roman';s.font.color.rgb=RGBColor(0,0,0)
  s.font.size=Pt(12);s.paragraph_format.space_after=Pt(0)
  s.paragraph_format.line_spacing=2
  if name not in ['Normal','Subtitle']:s.paragraph_format.keep_with_next=True
  rpr=s.element.get_or_add_rPr();rf=rpr.find(qn('w:rFonts'))
  if rf is not None:
   for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
    rf.attrib.pop(qn('w:'+attr),None)
   for attr in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+attr),'Times New Roman')
  ppr=s.element.get_or_add_pPr()
  for border in list(ppr.findall(qn('w:pBdr'))):ppr.remove(border)
  if name!='Normal':s.paragraph_format.first_line_indent=Inches(0)
 d.styles['Title'].font.size=Pt(18);d.styles['Title'].font.bold=True
 d.styles['Title'].paragraph_format.line_spacing=1.15
 d.styles['Title'].paragraph_format.space_after=Pt(12)
 d.styles['Normal'].paragraph_format.first_line_indent=Inches(.5)
 for name in ['Heading 1','Heading 2']:
  d.styles[name].font.bold=True;d.styles[name].paragraph_format.space_before=Pt(12)
  d.styles[name].paragraph_format.space_after=Pt(6);d.styles[name].paragraph_format.line_spacing=1.15
 footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
 lines=source.read_text(encoding='utf-8').splitlines();i=0;refs=False
 while i<len(lines):
  line=lines[i].strip()
  if not line:i+=1;continue
  if line.startswith('# '):
   p=d.add_paragraph(line[2:],'Title')
   ppr=p._p.get_or_add_pPr();border=OxmlElement('w:pBdr')
   for edge in ['top','bottom','left','right','between']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'nil');border.append(e)
   ppr.append(border)
   i+=1;continue
  if line.startswith('## '):
   h=line[3:];d.add_paragraph(h,'Heading 1');refs=h in ['References','Sources and reproducibility records'];i+=1;continue
  if line.startswith('|'):
   tablelines=[]
   while i<len(lines) and lines[i].strip().startswith('|'):
    s=lines[i].strip();i+=1
    if re.fullmatch(r'[| :\-]+',s):continue
    tablelines.append([v.strip() for v in s.strip('|').split('|')])
   t=d.add_table(rows=0,cols=len(tablelines[0]));t.autofit=False
   widths=([.6,3.3,1.2,1.4] if len(tablelines[0])==4 else [1.4,5.1])
   # Four-column model table has a longer contrast and short model label.
   if tablelines[0][0]!='Model' and len(widths)==4:widths=[2.6,1.3,1.2,1.4]
   for rownum,values in enumerate(tablelines):
    cells=t.add_row().cells
    for j,value in enumerate(values):
     cells[j].width=Inches(widths[j]);cells[j].text=value
     tcpr=cells[j]._tc.get_or_add_tcPr()
     margins=OxmlElement('w:tcMar')
     for edge in ['top','left','bottom','right']:
      e=OxmlElement('w:'+edge);e.set(qn('w:w'),'65');e.set(qn('w:type'),'dxa');margins.append(e)
     tcpr.append(margins)
     borders=OxmlElement('w:tcBorders')
     for edge in ['top','left','bottom','right']:
      e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
     tcpr.append(borders)
     if rownum==0:
      shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'E8E8E8');tcpr.append(shading)
     for p in cells[j].paragraphs:
      p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(3);p.paragraph_format.first_line_indent=Inches(0)
      if rownum==0:p.paragraph_format.keep_with_next=True
      for run in p.runs:run.font.size=Pt(11);run.bold=rownum==0
      if j>=2:p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    trpr=t.rows[-1]._tr.get_or_add_trPr();nosplit=OxmlElement('w:cantSplit');trpr.append(nosplit)
    if rownum==0:trpr.append(OxmlElement('w:tblHeader'))
   d.add_paragraph().paragraph_format.space_after=Pt(4)
   continue
  block=[line];i+=1
  while i<len(lines) and lines[i].strip() and not lines[i].startswith(('#','|')):
   block.append(lines[i].strip());i+=1
  p=d.add_paragraph(' '.join(block));p.paragraph_format.widow_control=True
  if refs and not p.text.startswith('COMD7290 assignment'):
   p.paragraph_format.left_indent=Inches(.5);p.paragraph_format.first_line_indent=Inches(-.5)
   p.paragraph_format.line_spacing=1.5
  if p.text.startswith('Prepared for'):
   p.paragraph_format.first_line_indent=Inches(0);p.paragraph_format.space_after=Pt(12)
 d.core_properties.title=lines[0].removeprefix('# ')
 d.core_properties.author='Ma Jiayi project with AI assistance'
 d.core_properties.subject='COMD7290 computational reproduction verification'
 target=source.with_suffix('.docx');d.save(target);print(target)
if __name__=='__main__':
 for name in ['01_results_comparison','02_python_stata_equivalence']:
  build(ROOT/'report/verification'/f'{name}.md')
